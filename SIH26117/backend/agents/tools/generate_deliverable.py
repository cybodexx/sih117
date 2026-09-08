"""generate_deliverable tool — writes a real Word/Excel/PowerPoint file to the vault.

Runs ONLY after a human operator approves at the HITL gate. Content is rendered
deterministically from the structured spec produced by the drafting step, so the
bytes on disk are exactly what the answer claims. Registry row + SHA-256 + audit
entry are written alongside the file.
"""
from __future__ import annotations

import hashlib
import io
import re
import uuid
from datetime import datetime, timezone

import structlog
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession

from backend.core.config import get_settings
from backend.db.models.deliverable import DeliverableModel
from backend.services.audit.writer import emit as audit_emit

logger = structlog.get_logger()
settings = get_settings()

KINDS = ("word", "excel", "powerpoint")
MIME = {
    "word": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "powerpoint": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
EXT = {"word": ".docx", "excel": ".xlsx", "powerpoint": ".pptx"}


class DeliverableSection(BaseModel):
    heading: str = ""
    paragraphs: list[str] = []
    bullets: list[str] = []
    rows: list[list[str]] = []


class DeliverableWord(BaseModel):
    title: str = ""
    sections: list[DeliverableSection] = []


class DeliverableExcel(BaseModel):
    title: str = ""
    sheet: str = "Data"
    rows: list[list[str]] = []


class DeliverableSlide(BaseModel):
    heading: str = ""
    bullets: list[str] = []


class DeliverablePptx(BaseModel):
    title: str = ""
    slides: list[DeliverableSlide] = []


DELIVERABLE_SCHEMAS: dict[str, type[BaseModel]] = {
    "word": DeliverableWord,
    "excel": DeliverableExcel,
    "powerpoint": DeliverablePptx,
}


def guess_kind(question: str) -> str:
    """Pick the deliverable file kind from the phrasing (default word)."""
    q = question.lower()
    if any(k in q for k in ("excel", "xlsx", "spreadsheet", "workbook", "tabulat", "work sheet", "data file")):
        return "excel"
    if any(k in q for k in ("powerpoint", "presentation", "pptx", "slide", "ppt deck", "deck")):
        return "powerpoint"
    return "word"


def _slug(text: str, max_len: int = 48) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "-", (text or "").strip()).strip("-").lower()
    return slug[:max_len].rstrip("-") or "deliverable"


def normalize_content(kind: str, raw: dict) -> dict:
    """Shape a drafting-schema dict into the renderer contract."""
    if kind == "excel":
        return {
            "title": raw.get("title") or "",
            "sheet": (raw.get("sheet") or "Data")[:31],
            "rows": [[str(c) if c is not None else "" for c in r] for r in (raw.get("rows") or [])],
        }
    out = {
        "title": (raw.get("title") or "")[:200],
        "sections": list(raw.get("sections") or []),
        "slides": list(raw.get("slides") or []),
    }
    return {k: v for k, v in out.items() if v or k == "title"}


def content_from_markdown(kind: str, title: str, md: str) -> dict:
    """Fallback structure builder when the model refuses structured JSON.

    Parses a plain markdown draft into the noise-free renderer contract, so a
    deliverable is still produced (honestly labelled as a fallback).
    """
    lines = [ln for ln in (md or "").splitlines() if ln.strip()]
    if kind == "word":
        sections: list[dict] = []
        cur: dict = {"heading": "", "paragraphs": [], "bullets": [], "rows": []}
        for ln in lines:
            if re.match(r"^#{1,6}\s+", ln):
                if cur["paragraphs"] or cur["bullets"]:
                    sections.append(cur)
                cur = {"heading": re.sub(r"^#{1,6}\s+", "", ln).strip(), "paragraphs": [], "bullets": [], "rows": []}
            elif re.match(r"^[\-*•]\s+", ln):
                cur["bullets"].append(re.sub(r"^[\-*•]\s+", "", ln).strip())
            else:
                cur["paragraphs"].append(ln.strip())
        if cur["paragraphs"] or cur["bullets"] or cur["heading"]:
            sections.append(cur)
        return {"title": title, "sections": sections}
    if kind == "excel":
        return {
            "title": title,
            "sheet": "Draft",
            "rows": [["Line", "Text"]] + [[str(i), ln] for i, ln in enumerate(lines, start=1)],
        }
    slides: list[dict] = []
    cur: dict = {"heading": "", "bullets": []}
    for ln in lines:
        if re.match(r"^#{1,6}\s+", ln):
            if cur["bullets"]:
                slides.append(cur)
            cur = {"heading": re.sub(r"^#{1,6}\s+", "", ln).strip(), "bullets": []}
        else:
            cur["bullets"].append(re.sub(r"^[\-*•]\s+", "", ln).strip())
    if cur["bullets"] or cur["heading"]:
        slides.append(cur)
    if not slides:
        slides.append({"heading": title, "bullets": lines})
    return {"title": title, "slides": slides}


def render_word(title: str, sections: list[dict]) -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading(title or "Deliverable", level=0)
    for sec in sections:
        if sec.get("heading"):
            doc.add_heading(sec["heading"], level=1)
        for p in sec.get("paragraphs") or []:
            doc.add_paragraph(p)
        for b in sec.get("bullets") or []:
            doc.add_paragraph(b, style="List Bullet")
        rows = sec.get("rows") or []
        if rows:
            table = doc.add_table(rows=1, cols=len(rows[0]))
            try:
                table.style = "Table Grid"
            except Exception:
                pass
            for j, cell in enumerate(rows[0]):
                table.rows[0].cells[j].text = str(cell)
            for r in rows[1:]:
                cells = table.add_row().cells
                for j in range(min(len(r), len(rows[0]))):
                    cells[j].text = str(r[j])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_excel(content: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = (content.get("sheet") or "Data")[:31]
    rows = content.get("rows") or []
    for ri, r in enumerate(rows):
        ws.append([str(c) if c is not None else "" for c in r])
        if ri == 0:
            for cell in ws[1]:
                cell.font = Font(bold=True)
    if rows:
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_pptx(content: dict) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    title = content.get("title") or "Deliverable"
    slides = content.get("slides") or []
    if not slides:
        slides = [{"heading": title, "bullets": []}]
    for i, sl in enumerate(slides):
        heading = sl.get("heading") or title
        if i == 0 and (heading == title or not heading) and slides and len(slides) > 1:
            layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = title
            continue
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = heading
        body = slide.placeholders[1]
        tf = body.text_frame
        for j, b in enumerate(sl.get("bullets") or []):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = b
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def render_bytes(kind: str, content: dict) -> bytes:
    if kind == "word":
        return render_word(content.get("title") or "", content.get("sections") or [])
    if kind == "excel":
        return render_excel(content)
    return render_pptx(content)


async def generate_deliverable(
    db: AsyncSession,
    *,
    user_id: str,
    session_id: str,
    turn_id: str,
    kind: str,
    title: str,
    content: dict,
) -> dict:
    """Render content to real file bytes, persist them and a registry row.

    Returns the file metadata (id, filename, mime, size, sha256, url).
    """
    kind = (kind or "word").lower()
    if kind not in KINDS:
        raise ValueError(f"unsupported deliverable kind: {kind}")
    content = normalize_content(kind, content)

    payload = render_bytes(kind, content)
    sha256 = hashlib.sha256(payload).hexdigest()

    vault_dir = settings.vault_path / "deliverables"
    vault_dir.mkdir(parents=True, exist_ok=True)
    if not str(vault_dir.resolve()).startswith(str(settings.vault_path.resolve())):
        raise RuntimeError("refusing to write deliverable outside the vault path")

    stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    filename = f"{_slug(title)}{EXT[kind]}"
    path = vault_dir / f"{stamp}_{uuid.uuid4().hex[:8]}_{filename}"
    path.write_bytes(payload)

    deliverable = DeliverableModel(
        id=uuid.uuid4(),
        user_id=uuid.UUID(user_id),
        session_id=uuid.UUID(session_id),
        turn_id=uuid.UUID(turn_id),
        kind=kind,
        title=(title or "")[:500],
        filename=filename,
        mime=MIME[kind],
        size_bytes=len(payload),
        storage_key=str(path),
        sha256=sha256,
        content_meta={
            "title": (title or "")[:200],
            "kind": kind,
            **({"sections": len(content.get("sections") or [])} if kind == "word" else {}),
            **({"rows": len(content.get("rows") or [])} if kind == "excel" else {}),
            **({"slides": len(content.get("slides") or [])} if kind == "powerpoint" else {}),
        },
        created_at=datetime.now(timezone.utc),
    )
    db.add(deliverable)
    await db.flush()
    await db.commit()

    result = {
        "deliverable_id": str(deliverable.id),
        "filename": filename,
        "kind": kind,
        "title": (title or "")[:500],
        "mime": MIME[kind],
        "size_bytes": len(payload),
        "sha256": sha256,
        "url": f"/api/v1/deliverables/{deliverable.id}/file",
    }

    await audit_emit(
        "DELIVERABLE_CREATED",
        correlation_id="",
        user_id=user_id,
        role="OPERATOR",
        resource_type="deliverable",
        resource_id=str(deliverable.id),
        document_ids=[],
        decision=json_dumps_safe(result),
        severity="info",
    )

    logger.info(
        "deliverable_created",
        deliverable_id=str(deliverable.id),
        kind=kind,
        filename=filename,
        size_bytes=len(payload),
        sha256=sha256[:16],
    )
    return result


def json_dumps_safe(obj: dict) -> str:
    import json

    try:
        return json.dumps(obj, sort_keys=True)
    except Exception:
        return json.dumps({"deliverable_id": obj.get("deliverable_id"), "size_bytes": obj.get("size_bytes")}, sort_keys=True)