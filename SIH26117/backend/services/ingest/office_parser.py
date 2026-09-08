"""Office document parsers: DOCX, PPTX, XLSX -> ingest Elements.

Uses python-docx / python-pptx / openpyxl. Excel sheets are also exported as
CSV text so the pipeline can materialise them into ds_* SQL tables.
"""
from __future__ import annotations

import csv
import io

from backend.services.ingest.pdf_parser import Element


def parse_docx(raw: bytes) -> list[Element]:
    """Parse a .docx into Elements, preserving paragraph/table reading order."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(raw))
    elements: list[Element] = []
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, doc).text.strip()
            if text:
                elements.append(Element(kind="TEXT", page=1, text=text))
        elif child.tag == qn("w:tbl"):
            rows: list[str] = []
            for row in DocxTable(child, doc).rows:
                rows.append(" | ".join(c.text.strip() for c in row.cells))
            if rows:
                elements.append(Element(kind="TABLE", page=1, text="\n".join(rows)))
    return elements


def parse_pptx(raw: bytes) -> list[Element]:
    """Parse a .pptx into one TEXT element per slide with text frames + tables."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(raw))
    elements: list[Element] = []
    for idx, slide in enumerate(prs.slides, start=1):
        blocks: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        blocks.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
        if blocks:
            elements.append(Element(kind="TEXT", page=idx, text="\n".join(blocks)))
    return elements


def xlsx_to_csv_text(raw: bytes) -> str:
    """Export the first worksheet of an .xlsx as CSV text (for SQL materialisation)."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        out = io.StringIO()
        writer = csv.writer(out)
        for row in ws.iter_rows(values_only=True):
            writer.writerow(["" if cell is None else str(cell) for cell in row])
        return out.getvalue()
    finally:
        wb.close()


def parse_xlsx(raw: bytes) -> tuple[list[Element], str]:
    """Parse an .xlsx into DATA_SLICE Elements (for RAG) plus CSV text (for SQL)."""
    csv_text = xlsx_to_csv_text(raw)
    lines = [ln for ln in csv_text.splitlines() if ln.strip()]
    elements: list[Element] = []
    if lines:
        header = lines[0]
        for i in range(1, len(lines) + 1, 200):
            block = "\n".join([header] + lines[i : i + 200])
            elements.append(Element(kind="DATA_SLICE", page=1, text=block))
    return elements, csv_text